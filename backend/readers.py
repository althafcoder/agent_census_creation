"""Read every supported upload into a common RawDoc structure.

Supported: .xlsx .xlsm .xls .csv .pdf (digital + scanned via OCR) .docx
"""
from __future__ import annotations

import importlib.util
import logging
import os
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd

log = logging.getLogger("census.readers")

SUPPORTED = {".xlsx", ".xlsm", ".xls", ".csv", ".pdf", ".docx", ".txt"}


@dataclass
class RawTable:
    name: str                     # sheet / table label
    df: pd.DataFrame              # header=None style grid (strings / objects)


@dataclass
class RawDoc:
    path: Path
    ext: str
    tables: list[RawTable] = field(default_factory=list)
    text: str = ""
    scanned: bool = False         # True when OCR was needed
    extraction_method: str = ""   # Engine or approach used
    extracted_text_path: Path | None = None
    form_fields: dict = field(default_factory=dict)   # Interactive PDF form fields (pypdf)
    conflicts: list[str] = field(default_factory=list) # Conflicts between text and form values

    @property
    def name(self) -> str:
        return self.path.name


# ------------------------------------------------------------------ excel ----
def _patch_openpyxl_cellrange():
    """Ensure openpyxl can handle column-only (e.g. 'P:Q') or row-only ('1:10') ranges.

    Some spreadsheets (e.g. LibreOffice or third-party exports) write whole-column
    data validations or conditional formatting like sqref='P:Q N9:N615'.
    Standard openpyxl crashes on these with:
    TypeError: expected <class 'openpyxl.worksheet.cell_range.MultiCellRange'>.
    """
    try:
        import openpyxl.worksheet.cell_range as cr
        from openpyxl.utils.cell import range_to_tuple, range_boundaries

        if getattr(cr, "_census_patched", False):
            return

        def safe_cell_range_init(self, range_string=None, min_col=None, min_row=None,
                                 max_col=None, max_row=None, title=None):
            if range_string is not None:
                if "!" in range_string:
                    title, (min_col, min_row, max_col, max_row) = range_to_tuple(range_string)
                else:
                    min_col, min_row, max_col, max_row = range_boundaries(range_string)
            if min_col is None:
                min_col = 1
            if max_col is None:
                max_col = 16384
            if min_row is None:
                min_row = 1
            if max_row is None:
                max_row = 1048576
            self.min_col = min_col
            self.min_row = min_row
            self.max_col = max_col
            self.max_row = max_row
            self.title = title

        cr.CellRange.__init__ = safe_cell_range_init
        cr._census_patched = True
    except Exception:
        pass


_patch_openpyxl_cellrange()


def _sanitize_excel_xml(src_path: Path) -> Path | None:
    """Create a sanitized temporary copy of an .xlsx/.xlsm archive with malformed XML stripped."""
    import re
    import zipfile
    try:
        fd, tmp_name = tempfile.mkstemp(suffix=src_path.suffix.lower())
        os.close(fd)
        tmp_path = Path(tmp_name)
        with zipfile.ZipFile(src_path, "r") as zin, zipfile.ZipFile(tmp_path, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                content = zin.read(item.filename)
                if item.filename.startswith("xl/worksheets/sheet") and item.filename.endswith(".xml"):
                    text = content.decode("utf-8", errors="ignore")
                    text = re.sub(r"<dataValidations[^>]*>.*?</dataValidations>", "", text, flags=re.DOTALL)
                    text = re.sub(r"<dataValidation[^>]*>.*?</dataValidation>", "", text, flags=re.DOTALL)
                    text = re.sub(r"<conditionalFormatting[^>]*>.*?</conditionalFormatting>", "", text, flags=re.DOTALL)
                    content = text.encode("utf-8")
                zout.writestr(item, content)
        return tmp_path
    except Exception as exc:
        log.warning("XML sanitization failed for %s: %s", src_path.name, exc)
        return None


def _xls_to_xlsx(path: Path) -> Path:
    """Convert legacy .xls with LibreOffice when xlrd is not installed."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError(f"Cannot read {path.name}: install xlrd or LibreOffice")
    out = Path(tempfile.mkdtemp())
    subprocess.run([soffice, "--headless", "--convert-to", "xlsx", "--outdir", str(out), str(path)],
                   check=True, capture_output=True, timeout=180)
    return out / (path.stem + ".xlsx")


def _read_excel(path: Path) -> RawDoc:
    doc = RawDoc(path, path.suffix.lower())
    doc.extraction_method = "Excel Spreadsheet (Direct Table Parsing)"
    target = path
    if doc.ext == ".xls":
        try:
            import xlrd  # noqa: F401
        except ImportError:
            target = _xls_to_xlsx(path)

    sheets = None
    # 1. Try calamine engine first if available (faster, resilient to openpyxl XML edge-cases)
    try:
        sheets = pd.read_excel(target, sheet_name=None, header=None, dtype=object, engine="calamine")
    except Exception:
        pass

    # 2. Standard pd.read_excel (openpyxl with patched CellRange)
    if sheets is None:
        try:
            sheets = pd.read_excel(target, sheet_name=None, header=None, dtype=object)
        except Exception as exc:
            # 3. Fallback: sanitize worksheet XML (strip malformed dataValidations / conditionalFormatting)
            sanitized = _sanitize_excel_xml(target)
            if sanitized:
                try:
                    sheets = pd.read_excel(sanitized, sheet_name=None, header=None, dtype=object)
                except Exception:
                    pass
                finally:
                    sanitized.unlink(missing_ok=True)
            if sheets is None:
                raise exc

    for name, df in sheets.items():
        # Preserve spreadsheet column positions and blank cells:
        # Only drop entirely empty rows; do not drop blank columns that would shift column alignment.
        df = df.dropna(how="all").reset_index(drop=True)
        if not df.empty:
            doc.tables.append(RawTable(str(name), df))
            doc.text += "\n".join(" | ".join("" if pd.isna(c) else str(c) for c in row)
                                  for row in df.values.tolist()[:60]) + "\n"
    return doc


def _read_csv(path: Path) -> RawDoc:
    df = pd.read_csv(path, header=None, dtype=object, keep_default_na=False, encoding_errors="replace")
    df = df.replace("", pd.NA).dropna(how="all").dropna(axis=1, how="all").reset_index(drop=True)
    doc = RawDoc(path, ".csv", [RawTable(path.stem, df)],
                  "\n".join(" | ".join("" if pd.isna(c) else str(c) for c in r) for r in df.values.tolist()[:60]))
    doc.extraction_method = "CSV File (Structured Table Parsing)"
    return doc


# -------------------------------------------------------------------- pdf ----
_paddle_ocr_engine = None


def _paddle_available() -> bool:
    """True when this interpreter can run PaddleOCR in-process."""
    return (importlib.util.find_spec("paddle") is not None
            and importlib.util.find_spec("paddleocr") is not None)


def _ocr_pdf_external(path: Path) -> str:
    """Run PaddleOCR in a separate interpreter that has paddlepaddle installed.

    Keeps OCR working when the pipeline runs on a newer Python (e.g. 3.14) that paddlepaddle has no
    wheel for yet, while the project's `venv` (3.11) does. Override the interpreter with
    PADDLE_OCR_PYTHON. Returns "" when no capable interpreter is found.
    """
    if os.getenv("PADDLE_OCR_BRIDGE"):
        return ""
    from census_config import BASE_DIR

    override = os.getenv("PADDLE_OCR_PYTHON", "")
    if override and Path(override).exists():
        py = Path(override)
    else:
        cand = (BASE_DIR / "venv" / ("Scripts/python.exe" if os.name == "nt" else "bin/python"))
        if not cand.exists():
            return ""
        py = cand

    fd, out_name = tempfile.mkstemp(suffix=".txt")
    os.close(fd)
    out = Path(out_name)
    env = os.environ.copy()
    env["PADDLE_OCR_BRIDGE"] = "1"
    log.info("Delegating PaddleOCR for %s to %s", path.name, py)
    try:
        proc = subprocess.run(
            [str(py), str(Path(__file__).with_name("ocr_bridge.py")), str(path), str(out)],
            capture_output=True, text=True, timeout=1800, env=env)
        text = out.read_text(encoding="utf-8", errors="replace") if out.exists() else ""
        if text.strip():
            return text
        log.warning("External PaddleOCR produced no text for %s (rc=%s): %s",
                    path.name, proc.returncode, (proc.stderr or "")[-400:])
    except Exception as exc:  # pragma: no cover - environment dependent
        log.warning("External PaddleOCR failed on %s: %s", path.name, exc)
    finally:
        try:
            out.unlink(missing_ok=True)
        except OSError:
            pass
    return ""


def _get_paddle_ocr():
    global _paddle_ocr_engine
    if _paddle_ocr_engine is None:
        from paddleocr import PaddleOCR
        try:
            _paddle_ocr_engine = PaddleOCR(use_angle_cls=True, lang="en", show_log=False)
        except Exception:
            _paddle_ocr_engine = PaddleOCR(use_angle_cls=True, lang="en")
    return _paddle_ocr_engine


def _reconstruct_ocr_layout(boxes_and_texts: list) -> str:
    """Group detected OCR polygon boxes into spatial lines and columns to preserve original document layout."""
    if not boxes_and_texts:
        return ""

    import numpy as np

    items = []
    for it in boxes_and_texts:
        if not it or len(it) < 2 or not isinstance(it[1], (tuple, list)):
            continue
        box = it[0]
        txt = str(it[1][0]).strip()
        if not txt:
            continue
        xs = [p[0] for p in box]
        ys = [p[1] for p in box]
        x_min, x_max = min(xs), max(xs)
        y_min, y_max = min(ys), max(ys)
        items.append({
            "text": txt,
            "x_min": x_min,
            "x_max": x_max,
            "y_min": y_min,
            "y_max": y_max,
            "y_center": (y_min + y_max) / 2.0,
            "height": max(1.0, y_max - y_min),
            "width": max(1.0, x_max - x_min),
        })

    if not items:
        return ""

    char_widths = [it["width"] / max(1, len(it["text"])) for it in items]
    avg_char_w = float(np.median(char_widths)) if char_widths else 10.0

    # Cluster boxes into lines based on vertical center overlap
    items.sort(key=lambda x: x["y_center"])
    lines = []
    for it in items:
        matched = False
        for line in lines:
            if abs(it["y_center"] - line["y_center"]) < max(it["height"], line["height"]) * 0.55:
                line["items"].append(it)
                line["y_center"] = sum(x["y_center"] for x in line["items"]) / len(line["items"])
                line["height"] = max(line["height"], it["height"])
                matched = True
                break
        if not matched:
            lines.append({"y_center": it["y_center"], "height": it["height"], "items": [it]})

    lines.sort(key=lambda l: l["y_center"])

    out_lines = []
    for l in lines:
        row = sorted(l["items"], key=lambda x: x["x_min"])
        row_str = ""
        last_x = 0
        for idx, it in enumerate(row):
            if idx == 0:
                indent = max(0, int(it["x_min"] / avg_char_w))
                row_str += " " * min(indent, 80) + it["text"]
            else:
                gap = it["x_min"] - last_x
                if gap > avg_char_w * 1.5:
                    spaces = max(2, int(gap / avg_char_w))
                    row_str += " " * min(spaces, 50) + it["text"]
                else:
                    row_str += " " + it["text"]
            last_x = it["x_max"]
        if row_str.strip():
            out_lines.append(row_str)

    return "\n".join(out_lines)


def _ocr_pdf(path: Path) -> str:
    # If paddlepaddle cannot run in this interpreter, delegate to the project's OCR interpreter.
    if not _paddle_available():
        text = _ocr_pdf_external(path)
        if text:
            return text
    try:
        from pdf2image import convert_from_path
        import numpy as np
        from census_config import POPPLER_PATH
    except ImportError as exc:
        log.warning("%s looks scanned but pdf2image/numpy are missing: %s", path.name, exc)
        return ""

    poppler = POPPLER_PATH if POPPLER_PATH and Path(POPPLER_PATH).exists() else None
    try:
        pages = convert_from_path(str(path), dpi=200, poppler_path=poppler)
    except Exception as exc:
        log.warning("Could not convert %s to images for OCR: %s", path.name, exc)
        return ""

    log.info("Running PaddleOCR on %s (%d page(s))...", path.name, len(pages))
    # Primary OCR: PaddleOCR
    try:
        ocr = _get_paddle_ocr()
        page_texts = []
        for pno, p in enumerate(pages, 1):
            img_arr = np.array(p)
            res = ocr.ocr(img_arr, cls=True)
            boxes = res[0] if res else []

            # Auto-detect orientation: if vertical tall boxes outnumber horizontal wide boxes
            tall_count = sum(
                1 for it in boxes
                if (max(pt[1] for pt in it[0]) - min(pt[1] for pt in it[0])) > (max(pt[0] for pt in it[0]) - min(pt[0] for pt in it[0])) * 1.2
            )
            wide_count = sum(
                1 for it in boxes
                if (max(pt[0] for pt in it[0]) - min(pt[0] for pt in it[0])) > (max(pt[1] for pt in it[0]) - min(pt[1] for pt in it[0])) * 1.2
            )

            if tall_count > wide_count and tall_count >= 5:
                # Document page is rotated sideways (landscape in portrait container)
                rot270 = p.rotate(270, expand=True)
                rot_np = np.array(rot270)
                is_upside_down = False
                try:
                    dt_boxes = ocr.text_detector(rot_np)
                    test_boxes = dt_boxes[0][:8] if dt_boxes and len(dt_boxes[0]) > 0 else []
                    crops = []
                    for b in test_boxes:
                        xs = [int(pt[0]) for pt in b]
                        ys = [int(pt[1]) for pt in b]
                        x1, x2 = max(0, min(xs)), min(rot_np.shape[1], max(xs))
                        y1, y2 = max(0, min(ys)), min(rot_np.shape[0], max(ys))
                        if x2 > x1 and y2 > y1:
                            crops.append(rot_np[y1:y2, x1:x2])
                    if crops:
                        cls_res = ocr.text_classifier(crops)[1]
                        count_180 = sum(1 for label, score in cls_res if label == "180")
                        if count_180 > len(crops) / 2:
                            is_upside_down = True
                except Exception:
                    pass

                final_img = p.rotate(90, expand=True) if is_upside_down else rot270
                res = ocr.ocr(np.array(final_img), cls=True)
                boxes = res[0] if res else []

            # Reconstruct 2D spatial layout mirroring original PDF visual structure
            page_text = _reconstruct_ocr_layout(boxes)
            header_str = f"--- [Page {pno} of {len(pages)}] ---"
            page_texts.append(f"{header_str}\n{page_text}" if page_text else header_str)
            line_count = len(page_text.splitlines()) if page_text else 0
            log.info("PaddleOCR page %d/%d for %s: preserved %d layout lines", pno, len(pages), path.name, line_count)

        extracted = "\n\n".join(page_texts)
        log.info("PaddleOCR completed successfully for %s: %d total characters extracted with layout", path.name, len(extracted))
        return extracted
    except Exception as exc:
        log.warning("PaddleOCR failed on %s: %s", path.name, exc)
        text = _ocr_pdf_external(path)
        if text:
            return text

    # Fallback to pytesseract if available
    try:
        import pytesseract
        from census_config import TESSERACT_PATH
        if TESSERACT_PATH:
            tess_p = Path(TESSERACT_PATH)
            if tess_p.is_dir():
                tess_exe = tess_p / "tesseract.exe"
                if tess_exe.exists():
                    pytesseract.pytesseract.tesseract_cmd = str(tess_exe)
                else:
                    pytesseract.pytesseract.tesseract_cmd = str(tess_p / "tesseract")
            elif tess_p.exists():
                pytesseract.pytesseract.tesseract_cmd = str(tess_p)
        return "\n".join(pytesseract.image_to_string(p, config="--psm 6") for p in pages)
    except Exception as exc:
        log.warning("OCR fallback failed on %s: %s", path.name, exc)
        return ""


def _read_pdf(path: Path) -> RawDoc:
    import pdfplumber

    doc = RawDoc(path, ".pdf")
    texts: list[str] = []
    with pdfplumber.open(str(path)) as pdf:
        for pno, page in enumerate(pdf.pages, 1):
            t = page.extract_text(layout=True)
            if not t or not t.strip():
                t = page.extract_text() or ""
            texts.append(t)
            found = 0
            for ti, tbl in enumerate(page.extract_tables() or [], 1):          # ruled tables
                df = pd.DataFrame(tbl).replace("", pd.NA).dropna(how="all").dropna(axis=1, how="all")
                if len(df) >= 2:
                    doc.tables.append(RawTable(f"p{pno}_t{ti}", df.reset_index(drop=True)))
                    found += 1
            if not found:                                                       # borderless tables
                try:
                    ts = {"vertical_strategy": "text", "horizontal_strategy": "text", "snap_tolerance": 4}
                    for ti, tbl in enumerate(page.extract_tables(ts) or [], 1):
                        df = pd.DataFrame(tbl).replace("", pd.NA).dropna(how="all").dropna(axis=1, how="all")
                        if len(df) >= 3 and df.shape[1] >= 3:
                            doc.tables.append(RawTable(f"p{pno}_x{ti}", df.reset_index(drop=True)))
                except Exception as exc:  # pragma: no cover
                    log.debug("text-strategy table extraction failed on %s p%s: %s", path.name, pno, exc)

    # Interactive form field extraction via pypdf (AcroForm fields: text, checkboxes, radio buttons)
    form_lines: list[str] = []
    try:
        from pypdf import PdfReader
        pypdf_reader = PdfReader(str(path))
        fields = pypdf_reader.get_fields() or {}
        for f_name, f_obj in fields.items():
            if not f_obj:
                continue
            f_val = getattr(f_obj, "value", None)
            f_type = str(getattr(f_obj, "field_type", ""))
            
            # Button/Checkbox/Radio (/Btn)
            if f_type == "/Btn":
                if f_val in (None, "/Off", "Off", False, ""):
                    status = "unchecked"
                    disp_val = "Off / Unchecked"
                else:
                    status = "checked"
                    disp_val = f"Checked ({f_val})"
                doc.form_fields[f_name] = {
                    "type": "button",
                    "status": status,
                    "raw_value": f_val,
                    "value": disp_val,
                    "is_checked": (status == "checked")
                }
                form_lines.append(f"{f_name}: {disp_val}")
            else:
                s_val = str(f_val).strip() if f_val is not None else ""
                status = "populated" if (s_val and s_val.lower() != "none") else "blank"
                doc.form_fields[f_name] = {
                    "type": f_type,
                    "status": status,
                    "raw_value": f_val,
                    "value": s_val if status == "populated" else ""
                }
                if status == "populated":
                    form_lines.append(f"{f_name}: {s_val}")
    except Exception as exc:
        log.debug("pypdf interactive form extraction failed or not applicable on %s: %s", path.name, exc)

    page_text = "\n".join(texts)
    if form_lines:
        form_block = "\n\n--- [INTERACTIVE FORM FIELDS] ---\n" + "\n".join(form_lines)
        doc.text = (page_text + form_block).strip()
    else:
        doc.text = page_text

    # "File identifier": digital vs scanned
    # Check whether the combined text has sufficient character density
    total_len = len(doc.text.strip())
    if total_len < 40 * max(1, len(texts)) * 0.25 and not form_lines:
        doc.scanned = True
        log.info("%s identified as SCANNED -> using PaddleOCR", path.name)
        ocr_text = _ocr_pdf(path)
        doc.text = (ocr_text + ("\n\n" + "\n".join(form_lines) if form_lines else "")).strip()
        doc.extraction_method = ("PaddleOCR (Scanned PDF OCR)" if ocr_text.strip()
                                 else "OCR unavailable (no text extracted)")
    else:
        doc.scanned = False
        method_parts = ["Digital PDF (pdfplumber Direct Extraction)"]
        if form_lines:
            method_parts.append(f"AcroForm Fields (pypdf: {len(form_lines)} fields)")
        doc.extraction_method = " + ".join(method_parts)
    return doc


# ------------------------------------------------------------------- docx ----
def _read_docx(path: Path) -> RawDoc:
    import docx

    d = docx.Document(str(path))
    doc = RawDoc(path, ".docx")
    doc.extraction_method = "Word Document (.docx Text & Tables)"
    doc.text = "\n".join(p.text for p in d.paragraphs if p.text.strip())
    for i, t in enumerate(d.tables, 1):
        rows = [[c.text.strip() for c in r.cells] for r in t.rows]
        df = pd.DataFrame(rows).replace("", pd.NA).dropna(how="all").dropna(axis=1, how="all")
        if not df.empty:
            doc.tables.append(RawTable(f"table{i}", df.reset_index(drop=True)))
            doc.text += "\n" + "\n".join(" | ".join(map(str, r)) for r in rows)
    return doc


def read_file(path: Path) -> RawDoc:
    ext = path.suffix.lower()
    if ext in {".xlsx", ".xlsm", ".xls"}:
        return _read_excel(path)
    if ext == ".csv":
        return _read_csv(path)
    if ext == ".pdf":
        return _read_pdf(path)
    if ext == ".docx":
        return _read_docx(path)
    if ext == ".txt":
        doc = RawDoc(path, ext, [], path.read_text(errors="replace"))
        doc.extraction_method = "Plain Text (.txt Direct Read)"
        return doc
    raise ValueError(f"Unsupported file type: {path.name}")


def save_extracted_text(doc: RawDoc, dest_dir: Path | None = None) -> Path:
    """Save extracted text and table summaries for a document to a clean text file."""
    import census_config as C
    dirs_to_save = [C.EXTRACTED_TEXT_DIR]
    if dest_dir:
        dirs_to_save.append(Path(dest_dir))

    txt_filename = f"{doc.path.stem}_extracted.txt".replace(" ", "_")
    last_saved: Path | None = None

    lines = [
        "=" * 80,
        f"DOCUMENT: {doc.name}",
        f"SOURCE PATH: {doc.path}",
        f"FILE EXTENSION: {doc.ext}",
        f"IS SCANNED (OCR USED): {doc.scanned}",
        f"EXTRACTION METHOD: {doc.extraction_method}",
        f"TOTAL CHARACTERS: {len(doc.text):,} chars",
        f"TOTAL TABLES DETECTED: {len(doc.tables)} table(s)",
        "=" * 80,
        "",
        "--- [EXTRACTED TEXT CONTENT] ---",
        doc.text.strip() if doc.text.strip() else "[No plain text content extracted]",
        "",
    ]

    if doc.tables:
        lines.append("--- [EXTRACTED TABLES SUMMARY] ---")
        for idx, t in enumerate(doc.tables, 1):
            lines.append(f"\n[Table {idx}: {t.name} (Rows: {len(t.df)}, Columns: {t.df.shape[1]})]")
            try:
                tbl_sample = t.df.head(60).to_string(index=False, header=False)
                lines.append(tbl_sample)
            except Exception:
                lines.append("[Table string conversion omitted]")
            lines.append("-" * 60)

    content = "\n".join(lines)

    for d in dirs_to_save:
        d.mkdir(parents=True, exist_ok=True)
        target = d / txt_filename
        target.write_text(content, encoding="utf-8")
        last_saved = target

    doc.extracted_text_path = last_saved
    if last_saved:
        log.info("      [EXTRACT] Saved structured text -> %s (%d chars)", last_saved.name, len(doc.text))
    return last_saved or (C.EXTRACTED_TEXT_DIR / txt_filename)


def read_all(source: Path | list[Path | str], dest_dir: Path | None = None) -> list[RawDoc]:
    raw_paths = [Path(p) for p in source] if isinstance(source, list) else [Path(source)]
    paths: list[Path] = []
    for p in raw_paths:
        if p.is_dir():
            paths.extend(sorted(p.iterdir()))
        elif p.is_file():
            paths.append(p)

    docs = []
    for p in paths:
        if p.name.startswith(("~$", ".")) or p.suffix.lower() not in SUPPORTED:
            continue
        try:
            log.info("   -> Reading file: %-42s (%s)", p.name, p.suffix.lower())
            doc = read_file(p)
            save_extracted_text(doc, dest_dir)
            docs.append(doc)
        except Exception as exc:  # keep going - one bad file must not kill the run
            log.error("Could not read %s: %s", p.name, exc)
    return docs

