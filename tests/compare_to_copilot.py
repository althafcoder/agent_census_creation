"""Golden comparison harness: diff our Marsh census output against a Copilot reference workbook.

Usage:
    python tests/compare_to_copilot.py <copilot.xlsx> [ours.xlsx]
If ours is omitted, the newest output/Marsh_Census_*.xlsx is used.
"""
from __future__ import annotations

import sys
from collections import Counter
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent


def _norm(v) -> str:
    if v is None:
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%m/%d/%Y")
    s = str(v).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s.upper()


def _find_header(ws, max_scan: int = 30):
    for r in range(1, max_scan + 1):
        vals = [_norm(c.value) for c in ws[r]]
        if "FIRST NAME" in vals and "LAST NAME" in vals:
            return r, {v: i for i, v in enumerate(vals) if v}
    raise ValueError(f"No header row found in sheet {ws.title}")


def load_rows(path: Path) -> tuple[dict, dict, list[str]]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    hr, cols = _find_header(ws)
    rows = {}
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        first = _norm(row[cols["FIRST NAME"]])
        last = _norm(row[cols["LAST NAME"]])
        if not first and not last:
            continue
        rec = {h: _norm(row[i]) if i < len(row) else "" for h, i in cols.items()}
        key = (first, last, rec.get("DATE OF BIRTH", ""))
        n = 1
        while key in rows:
            n += 1
            key = (first, last, f"{rec.get('DATE OF BIRTH', '')}#{n}")
        rows[key] = rec
    return rows, cols, wb.sheetnames


def main() -> None:
    ref_path = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "Lamb Farms Inc - Marsh PEO Census Comparison.xlsx"
    if len(sys.argv) > 2:
        ours_path = Path(sys.argv[2])
    else:
        ours_path = max((ROOT / "output").glob("Marsh_Census_*.xlsx"), key=lambda p: p.stat().st_mtime)

    ref, ref_cols, ref_sheets = load_rows(ref_path)
    ours, our_cols, our_sheets = load_rows(ours_path)
    print(f"REFERENCE: {ref_path.name}  sheets={ref_sheets}  rows={len(ref)}")
    print(f"OURS     : {ours_path.name}  sheets={our_sheets}  rows={len(ours)}")
    print(f"Ref headers : {list(ref_cols)}")
    print(f"Our headers : {list(our_cols)}\n")

    only_ref = [k for k in ref if k not in ours]
    only_ours = [k for k in ours if k not in ref]
    print(f"Rows only in Copilot: {len(only_ref)}")
    for k in only_ref[:40]:
        print("   -", k, ref[k].get("MEMBER TYPE", ""), ref[k].get("MEDICAL COVERAGE", ""))
    print(f"Rows only in ours   : {len(only_ours)}")
    for k in only_ours[:40]:
        print("   +", k, ours[k].get("MEMBER TYPE", ""), ours[k].get("MEDICAL COVERAGE", ""))

    common = [k for k in ref if k in ours]
    shared_cols = [c for c in ref_cols if c in our_cols and c not in ("ROW #",)]
    diffs = Counter()
    examples: dict[str, list] = {}
    for k in common:
        for c in shared_cols:
            a, b = ref[k].get(c, ""), ours[k].get(c, "")
            if a != b:
                diffs[c] += 1
                examples.setdefault(c, []).append((k[0], k[1], a, b))

    total_cells = max(1, len(common) * len(shared_cols))
    mismatched = sum(diffs.values())
    print(f"\nMatched rows: {len(common)}  | Cell agreement: {100 * (1 - mismatched / total_cells):.1f}%")
    print("Per-column mismatches (Copilot -> Ours):")
    for c, n in diffs.most_common():
        print(f"  {c:<28} {n:>4}")
        for ex in examples[c][:6]:
            print(f"      {ex[0]} {ex[1]}: '{ex[2][:60]}' -> '{ex[3][:60]}'")

    for c in ("MEDICAL COVERAGE", "MEMBER TYPE"):
        if c in ref_cols and c in our_cols:
            print(f"\n{c} distribution  Copilot={dict(Counter(r[c] for r in ref.values()))}")
            print(f"{' ' * len(c)} distribution  Ours   ={dict(Counter(r[c] for r in ours.values()))}")


if __name__ == "__main__":
    main()
